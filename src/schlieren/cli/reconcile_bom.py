"""Compare an edited copy of the BOM workbook with bom/bom.csv, and optionally apply its edits to the CSV.

bom/bom.csv is authoritative. The workbook is a generated view, but it is the convenient place to edit the
BOM in a spreadsheet program. This command reads the workbook's BOM tab, matches rows to the CSV by ID, and
reports rows added, removed, and changed, field by field. Nothing is written unless --apply is given; review
the report (and afterwards `git diff bom/bom.csv`) before accepting the result.

    uv run reconcile-bom [workbook.xlsx] [--apply]

The exit status is 1 when differences were found and not applied. After --apply, run the tests and
`uv run build-bom-xlsx` to regenerate the workbook from the CSV.
"""

import argparse
import csv
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from schlieren.bom import ID_COLUMN, check_rows
from schlieren.cli.build_bom_xlsx import BOM_CSV, BOM_SHEET, DEFAULT_OUTPUT, read_bom

# Applying a workbook that drops more than this fraction of the BOM's rows (and more than a handful) needs
# --allow-many-removals: it is more likely a truncated sheet than an intended edit.
MANY_REMOVALS_FRACTION = 0.25
MANY_REMOVALS_MINIMUM = 5


@dataclass
class BomDiff:
    added: list[list[str]] = field(default_factory=list)
    removed: list[list[str]] = field(default_factory=list)
    changed: dict[str, list[tuple[str, str, str]]] = field(
        default_factory=dict
    )  # ID -> (column, CSV, workbook)
    reordered: bool = False

    def __bool__(self) -> bool:
        return bool(self.added or self.removed or self.changed or self.reordered)


def cell_text(value) -> str:
    """A workbook cell as the CSV would hold it; spreadsheet programs save whole numbers as floats.

    The BOM tab is data: a formula, a date, or any other type has no single faithful CSV text, so is refused.
    """
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, str | int | float):
        raise SystemExit(
            f"The BOM tab has a {type(value).__name__} cell ({value!r}); it may hold only text and numbers"
        )
    if isinstance(value, str) and value.startswith("="):
        raise SystemExit(
            f"The BOM tab has a formula ({value!r}); replace it with its value before reconciling"
        )
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_workbook_bom(path: Path) -> tuple[list[str], list[list[str]]]:
    # Formulas are read as formulas, not as cached values (which a program that never calculated the sheet
    # leaves empty), so that cell_text can refuse them.
    wb = load_workbook(path, read_only=True, data_only=False)
    if BOM_SHEET not in wb.sheetnames:
        raise SystemExit(f"{path} has no '{BOM_SHEET}' tab")
    table = [[cell_text(v) for v in row] for row in wb[BOM_SHEET].iter_rows(values_only=True)]
    wb.close()
    table = [row for row in table if any(row)]  # Spreadsheet programs may keep formatted but empty rows.
    if not table:
        raise SystemExit(f"The '{BOM_SHEET}' tab of {path} is empty")
    header = table[0]
    while header and header[-1] == "":
        header.pop()
    rows = []
    for row in table[1:]:
        if any(row[len(header) :]):
            raise SystemExit(f"Row {row[0] or '(no ID)'} has data beyond the '{header[-1]}' column")
        rows.append((row + [""] * len(header))[: len(header)])
    return header, rows


def rows_by_id(header: list[str], rows: list[list[str]], source: str) -> dict[str, list[str]]:
    id_index = header.index(ID_COLUMN)
    by_id = {}
    for row in rows:
        row_id = row[id_index]
        if not row_id:
            raise SystemExit(f"{source} has a row without an {ID_COLUMN}: {row}")
        if row_id in by_id:
            raise SystemExit(f"{source} has more than one row with {ID_COLUMN} {row_id}")
        by_id[row_id] = row
    return by_id


def diff_bom(header: list[str], csv_rows: list[list[str]], workbook_rows: list[list[str]]) -> BomDiff:
    old = rows_by_id(header, csv_rows, "bom.csv")
    new = rows_by_id(header, workbook_rows, "The workbook")
    diff = BomDiff(
        added=[row for row_id, row in new.items() if row_id not in old],
        removed=[row for row_id, row in old.items() if row_id not in new],
    )
    for row_id, row in new.items():
        fields = [(col, a, b) for col, a, b in zip(header, old.get(row_id, row), row, strict=True) if a != b]
        if fields:
            diff.changed[row_id] = fields
    diff.reordered = [i for i in old if i in new] != [i for i in new if i in old]
    return diff


def report(diff: BomDiff, header: list[str]) -> str:
    if not diff:
        return "The workbook matches bom/bom.csv."
    lines = []
    for title, rows in (("Added in the workbook", diff.added), ("Removed in the workbook", diff.removed)):
        lines.extend(
            f"{title}: " + "; ".join(f"{col}={val}" for col, val in zip(header, row, strict=True) if val)
            for row in rows
        )
    for row_id, fields in diff.changed.items():
        lines.append(f"Changed {row_id}:")
        lines.extend(f"    {col}: {a!r} -> {b!r}" for col, a, b in fields)
    if diff.reordered:
        lines.append("Row order differs from bom.csv.")
    counts = f"{len(diff.added)} added, {len(diff.removed)} removed, {len(diff.changed)} changed"
    return "\n".join([*lines, counts])


def write_bom(path: Path, header: list[str], rows: list[list[str]]) -> None:
    """Write the BOM as canonical CSV, replacing path atomically so an interruption cannot leave half a file."""
    fd, temporary = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8") as f:
            csv.writer(f, lineterminator="\n").writerows([header, *rows])
        os.replace(temporary, path)
    except BaseException:
        Path(temporary).unlink(missing_ok=True)
        raise


def too_many_removals(diff: BomDiff, row_count: int) -> bool:
    """Whether the workbook drops enough of the BOM that it should be confirmed rather than applied."""
    removed = len(diff.removed)
    return removed > MANY_REMOVALS_MINIMUM and removed > MANY_REMOVALS_FRACTION * row_count


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook", type=Path, nargs="?", default=DEFAULT_OUTPUT, help="Edited .xlsx")
    parser.add_argument("--apply", action="store_true", help="Write the workbook's BOM rows to bom/bom.csv")
    parser.add_argument(
        "--allow-many-removals",
        action="store_true",
        help=f"Apply even if the workbook drops more than {MANY_REMOVALS_FRACTION:.0%} of the rows",
    )
    args = parser.parse_args()
    header, csv_rows = read_bom(BOM_CSV)
    workbook_header, workbook_rows = read_workbook_bom(args.workbook)
    if workbook_header != header:
        raise SystemExit(f"Workbook columns {workbook_header} do not match bom.csv columns {header}")
    diff = diff_bom(header, csv_rows, workbook_rows)
    print(report(diff, header))
    if not diff:
        return
    if not args.apply:
        raise SystemExit(1)
    problems = check_rows(workbook_header, workbook_rows)
    if problems:
        raise SystemExit("Not applied; the workbook's BOM is not sound:\n  " + "\n  ".join(problems))
    if too_many_removals(diff, len(csv_rows)) and not args.allow_many_removals:
        raise SystemExit(
            f"Not applied: the workbook drops {len(diff.removed)} of {len(csv_rows)} rows. "
            "If that is intended, repeat with --allow-many-removals."
        )
    write_bom(BOM_CSV, header, workbook_rows)
    print(f"Wrote {BOM_CSV}; review `git diff`, run the tests, then `uv run build-bom-xlsx`.")


if __name__ == "__main__":
    main()
